# -*- coding: utf-8 -*-
"""
Étage 2 Pappers (étape 9 du backlog) : cibles, fusion des résultats, signal succession
====================================================================================
Pappers est payant (jetons par appel). Ce script ne fait AUCUN appel : il
prépare la liste des cibles, fusionne des résultats obtenus par ailleurs
(connecteur Pappers en session Claude, ou export manuel) et calcule le signal
de succession. Règle du backlog : nombre d'appels et coût annoncés à Ahmed
avant tout lot.

Étapes :
  python tools/enrich_pappers.py targets      -> tools/cache/pappers_targets.json
        (top 50 élec + top 30 zone test + top 50 injection, dédupliqués ; les
        sites déjà résolus en juin 2026 sont pré-remplis depuis
        enrichissement_pappers.json, OneDrive Screening)
  python tools/enrich_pappers.py merge        -> lit tools/cache/pappers_results.json
        (format ci-dessous), écrit `pappers` dans data/*.json, produit le
        classeur Nautilus_Qualification_Sites_v2_Pappers_<date>.xlsx (OneDrive)

Format d'un résultat (une entrée par clé de site) :
  {"cle": "17W…", "siren": "…", "exploitant": "…", "forme": "…", "naf": "…",
   "siege_commune": "…", "dirigeants": [{"nom","prenom","age","qualite"}],
   "beneficiaires": [{"nom","prenom","pourcentage","type"}],
   "resultat_net": 123, "annee_resultat": 2024, "fiabilite": "HAUTE|MOYENNE|FAIBLE|NON RESOLU",
   "note": "…", "source": "Pappers MCP 2026-09-…"}
"""
import json
import sys
import unicodedata
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
DATA = REPO / "data"
CACHE = REPO / "tools" / "cache"
JUNE = Path(r"C:\Users\ahmed\OneDrive\Bureau\__Nautilus\Nautilus Business\2__Pipeline\Active"
            r"\Biomethane platform\Biomethane France\Screening\enrichissement_pappers.json")
XLSX = Path(r"C:\Users\ahmed\OneDrive\Bureau\__Nautilus\Nautilus Business\2__Pipeline\Active"
            r"\Biomethane platform\Biomethane France\Screening") / f"Nautilus_Qualification_Sites_v2_Pappers_{date.today().isoformat()}.xlsx"
PARAMS = json.load(open(REPO / "tools" / "screening_params.json", encoding="utf-8"))
ZT = set(PARAMS["geographie"]["zone_test"])
NAF_AGRI = ("01.",)  # culture et production animale
NAF_ENERGIE = ("35.11Z", "35.21Z")


def norm(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return " ".join(s.lower().replace("-", " ").replace("'", " ").split())


def load_data():
    elec = json.load(open(DATA / "cogenerations.json", encoding="utf-8"))
    inj = json.load(open(DATA / "points-injection.json", encoding="utf-8"))
    return elec, inj


def targets():
    elec, inj = load_data()
    se = sorted([d for d in elec if d.get("score_v2") is not None], key=lambda d: -d["score_v2"])
    si = sorted([d for d in inj if d.get("score_v2") is not None], key=lambda d: -d["score_v2"])
    sel = {}
    for d in se[:50] + [d for d in se if d["region"] in ZT][:30]:
        sel[d["code_eic"]] = {"cle": d["code_eic"], "base": "elec", "nom": d["nom"], "commune": d["commune"],
                              "departement": d["departement"], "code_insee": d.get("code_insee"),
                              "puissance_kw": d.get("puissance_kw"), "score": d["score_v2"], "priorite": d["priorite_v2"],
                              "nom_masque": d["nom"] in ("Confidentiel", "-", "")}
    for d in si[:50]:
        sel[str(d["id_unique_projet"])] = {"cle": str(d["id_unique_projet"]), "base": "injection", "nom": d["nom_du_projet"],
                                           "commune": d["commune"], "departement": d["departement"],
                                           "code_insee": d.get("code_commune"), "capacite_gwh": d.get("capacite_de_production_gwh_an"),
                                           "score": d["score_v2"], "priorite": d["priorite_v2"], "nom_masque": False}
    # pré-remplissage depuis juin 2026 (mêmes noms de site)
    june = {norm(s["site"]): s for s in json.load(open(JUNE, encoding="utf-8"))["sites"]} if JUNE.exists() else {}
    n_june = 0
    for t in sel.values():
        j = june.get(norm(t["nom"])) or june.get(norm("Confidentiel|" + t["commune"]))
        if j and j.get("fiabilite") != "NON RESOLU":
            t["juin_2026"] = j
            n_june += 1
    CACHE.mkdir(parents=True, exist_ok=True)
    out = CACHE / "pappers_targets.json"
    json.dump(list(sel.values()), open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    masked = sum(1 for t in sel.values() if t["nom_masque"])
    print(f"Cibles : {len(sel)} sites ({sum(1 for t in sel.values() if t['base']=='elec')} élec, "
          f"{sum(1 for t in sel.values() if t['base']=='injection')} injection) ; nom masqué « Confidentiel » : {masked} ; "
          f"déjà résolus en juin 2026 : {n_june} ; à rechercher : {len(sel) - n_june} (dont {masked} par commune)")
    print("Fichier :", out)


def signal_succession(dirs):
    if dirs is None:
        return "À compléter"
    ages = [x.get("age") for x in dirs if x.get("age")]
    if not ages:
        return "Groupe / personnes morales"
    mx, mn = max(ages), min(ages)
    if mx >= 62 and mn >= 50:
        return f"FORT (dirigeant {mx} ans, pas de relève < 50)"
    if mx >= 62:
        return f"Transition en cours ({mx} ans / relève {mn} ans)"
    if mx >= 55:
        return f"Modéré (aîné {mx} ans)"
    return f"Faible (aîné {mx} ans)"


def capital_agricole(r):
    """Lecture prudente : personnes physiques + NAF agricole ou société civile
    d'exploitation = probablement agricole ; groupe / NAF énergie ou déchets
    avec bénéficiaire personne morale = probablement non agricole."""
    naf = (r.get("naf") or "")
    forme = (r.get("forme") or "").lower()
    ben = r.get("beneficiaires") or []
    pm = any((b.get("type") or "").lower().startswith("personne morale") for b in ben)
    if naf.startswith(NAF_AGRI) or "gaec" in forme or "earl" in forme or "scea" in forme:
        return "probablement ≥ 51 % agricole"
    if pm and naf.startswith(NAF_ENERGIE + ("38.", "37.")):
        return "probablement < 51 % agricole (groupe)"
    if ben and not pm:
        return "personnes physiques (agricole à confirmer)"
    return "inconnue"


def merge():
    elec, inj = load_data()
    res_path = CACHE / "pappers_results.json"
    results = {r["cle"]: r for r in json.load(open(res_path, encoding="utf-8"))} if res_path.exists() else {}
    tg = {t["cle"]: t for t in json.load(open(CACHE / "pappers_targets.json", encoding="utf-8"))}
    # juin 2026 en secours
    for cle, t in tg.items():
        if cle not in results and t.get("juin_2026"):
            j = t["juin_2026"]
            results[cle] = {"cle": cle, "siren": j.get("siren"), "exploitant": j.get("exploitant"), "forme": j.get("forme"),
                            "naf": j.get("naf"), "dirigeants": j.get("dirigeants"), "resultat_net": j.get("resultat"),
                            "annee_resultat": j.get("annee_resultat"), "fiabilite": j.get("fiabilite"),
                            "note": j.get("note"), "source": "Pappers MCP 2026-06-10"}
    n = 0
    for d in elec + inj:
        cle = d.get("code_eic") if "code_eic" in d else str(d.get("id_unique_projet"))
        r = results.get(cle)
        if not r:
            d.pop("pappers", None)
            continue
        d["pappers"] = {
            "siren": r.get("siren"), "exploitant": r.get("exploitant"), "forme": r.get("forme"), "naf": r.get("naf"),
            "siege_commune": r.get("siege_commune"),
            "dirigeants": r.get("dirigeants"), "signal_succession": signal_succession(r.get("dirigeants")) if r.get("fiabilite") != "NON RESOLU" else "Non résolu",
            "beneficiaires": r.get("beneficiaires"), "capital_agricole": capital_agricole(r),
            "resultat_net": r.get("resultat_net"), "annee_resultat": r.get("annee_resultat"),
            "fiabilite": r.get("fiabilite"), "note": r.get("note"), "source": r.get("source"),
        }
        n += 1
    json.dump(elec, open(DATA / "cogenerations.json", "w", encoding="utf-8"), ensure_ascii=False)
    json.dump(inj, open(DATA / "points-injection.json", "w", encoding="utf-8"), ensure_ascii=False)
    print(f"Fusion : {n} sites enrichis ({len(results)} résultats, {len(tg)} cibles)")

    # ---- classeur
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Étage 2 Pappers"
    headers = ["Base", "Priorité", "Score", "Clé", "Nom du site", "Commune", "Département", "Puissance kWé / Capacité GWh",
               "SIREN", "Exploitant", "Forme", "NAF", "Siège (commune)", "Dirigeants (âge, qualité)", "Signal succession",
               "Bénéficiaires effectifs", "Capital agricole (lecture prudente)", "Résultat net (€)", "Année", "Fiabilité", "Notes", "Source",
               "Contact", "Statut prospection", "Commentaires AdlF"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", name="Arial", size=9)
        c.fill = PatternFill("solid", fgColor="1E4260")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    rows = []
    for d in elec + inj:
        p = d.get("pappers")
        if not p:
            continue
        is_e = "code_eic" in d
        dirs = " ; ".join(f"{x.get('nom','')} {x.get('prenom') or ''} ({x.get('age') or 'PM'}, {x.get('qualite','')})".replace("  ", " ")
                          for x in (p.get("dirigeants") or [])) or ""
        ben = " ; ".join(f"{b.get('nom','')} {b.get('prenom') or ''} {b.get('pourcentage') or ''}%".strip()
                         for b in (p.get("beneficiaires") or [])) or ""
        rows.append(["Élec. biogaz" if is_e else "Injection", d.get("priorite_v2"), d.get("score_v2"),
                     d.get("code_eic") if is_e else d.get("id_unique_projet"),
                     d.get("nom") if is_e else d.get("nom_du_projet"), d.get("commune"), d.get("departement"),
                     d.get("puissance_kw") if is_e else d.get("capacite_de_production_gwh_an"),
                     p.get("siren"), p.get("exploitant"), p.get("forme"), p.get("naf"), p.get("siege_commune"), dirs,
                     p.get("signal_succession"), ben, p.get("capital_agricole"), p.get("resultat_net"), p.get("annee_resultat"),
                     p.get("fiabilite"), p.get("note"), p.get("source"), "", "", ""])
    order = {"FORT": 0, "Transition": 1, "Modéré": 2, "Groupe": 3, "Faible": 4, "À compléter": 5, "Non résolu": 6}
    rows.sort(key=lambda r: (order.get(str(r[14]).split(" ")[0], 7), -(r[2] or 0)))
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=9)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = 14 if len(h) < 14 else min(44, len(h) + 4)
    wb.save(XLSX)
    print("Classeur :", XLSX, f"({len(rows)} lignes)")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    cmd = sys.argv[1] if len(sys.argv) > 1 else "targets"
    {"targets": targets, "merge": merge}[cmd]()
