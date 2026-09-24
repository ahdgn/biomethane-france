# -*- coding: utf-8 -*-
"""
Score v2 et shortlist (étape 7 du backlog)
==========================================
Calcule un score sur 100 pour les sites du périmètre prospection v2, l'écrit
dans les données de l'app (`score_v2`, `priorite_v2`, `score_detail`) et
produit le classeur Excel de travail pour Antoine de la Faire.

Pondérations et seuils de priorité : tools/screening_params.json (`score_v2`).
Les grilles (fraction du poids par critère) sont ci-dessous et recopiées dans
la feuille « Méthodologie » du classeur. Statut : proposition, à valider avec
AdlF ; changer une pondération = éditer le JSON et relancer.

Périmètres scorés (les autres sites gardent un score vide) :
  · électricité biogaz : code combustible B.MET, en service, ≥ plancher kWé ;
  · injection : types agricoles + industriel territorial, site ouvert, ≥ 5 GWh/an.

Usage : python tools/qualify_v2.py [--xlsx chemin_de_sortie.xlsx]
"""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent.parent
DATA = REPO / "data"
PARAMS = json.load(open(REPO / "tools" / "screening_params.json", encoding="utf-8"))
ANNEE = date.today().year
DEFAULT_XLSX = (r"C:\Users\ahmed\OneDrive\Bureau\__Nautilus\Nautilus Business\2__Pipeline\Active"
                r"\Biomethane platform\Biomethane France\Screening"
                rf"\Nautilus_Qualification_Sites_v2_{date.today().isoformat()}.xlsx")

W = PARAMS["score_v2"]
PRIO = W["priorites"]

# ------------------------------------------------------------------ grilles
GRILLES = {
    "elec": {
        "puissance": "≥ 1 000 kWé : 100 % · 500-999 : 80 % · 250-499 : 50 % (BC 14/09 : plancher 250, cible 500, priorité 1 MW)",
        "fenetre_echeance": "échéance estimée ≤ 2028 : 100 % · 2029-2030 : 80 % · 2031-2033 : 50 % · 2034 et après : 25 % · inconnue : 0 (BC : « dans 1, 2, 5, 10 ans »)",
        "coefficient_cpb": "coefficient CPB estimé à l'année de conversion par défaut : 1 → 100 % · 0,95 → 80 % · 0,8 → 0 (arrêté du 26/12/2025)",
        "reseau": "distance GRDF ≤ 2 km : 100 % · ≤ 4 : 80 % · ≤ 5 : 60 % · ≤ 10 : 25 % · au-delà ou inconnue : 0 (GRDF 4 km, BC/AdlF 4-5 km)",
        "facteur_charge": "énergie injectée / (puissance × 8 760 h) ≥ 0,7 : 100 % · 0,5-0,7 : 60 % · < 0,5 : 30 % · inconnu : 40 % (proxy de qualité d'exploitation, en attendant les intrants du registre équipe)",
    },
    "injection": {
        "capacite": "10-25 GWh/an : 100 % · 25-50 : 80 % · 5-10 : 60 % · > 50 : 40 % (enveloppe 5-7 M€ par objet, BC)",
        "tarif_residuel": "années de tarif restantes (MES + 15 − année courante) : 6-10 : 100 % · 11-12 : 70 % · 3-5 : 50 % · ≥ 13 : 40 % · < 3 : 30 % (weekly 26/03 : « de bons actifs avec 7-8 ans de tarif » ; < 3 ans = passerelle CPB avant fin 2027)",
        "type": "agricole autonome : 100 % · agricole territorial : 80 % · industriel territorial : 50 %",
        "augmentation_prevue": "augmentation de capacité prévue au registre : 100 % · aucune : 40 % (site en développement = exploitant investisseur, potentiel d'extension)",
    },
}


def frac_puissance(kw):
    if kw is None:
        return 0
    return 1.0 if kw >= 1000 else 0.8 if kw >= 500 else 0.5 if kw >= 250 else 0


def frac_echeance(e):
    if e is None:
        return 0
    return 1.0 if e <= 2028 else 0.8 if e <= 2030 else 0.5 if e <= 2033 else 0.25


def frac_coef(coef):
    return {1: 1.0, 0.95: 0.8, 0.8: 0}.get(coef, 0)


def frac_reseau(km):
    if km is None:
        return 0
    return 1.0 if km <= 2 else 0.8 if km <= 4 else 0.6 if km <= 5 else 0.25 if km <= 10 else 0


def frac_charge(fc):
    if fc is None:
        return 0.4
    return 1.0 if fc >= 0.7 else 0.6 if fc >= 0.5 else 0.3


def frac_capacite(gwh):
    if gwh is None:
        return 0
    return 1.0 if 10 <= gwh <= 25 else 0.8 if 25 < gwh <= 50 else 0.6 if 5 <= gwh < 10 else 0.4 if gwh > 50 else 0


def frac_tarif(rest):
    if rest is None:
        return 0
    return 1.0 if 6 <= rest <= 10 else 0.7 if 11 <= rest <= 12 else 0.5 if 3 <= rest <= 5 else 0.4 if rest >= 13 else 0.3


def frac_type(t):
    return {"Agricole autonome": 1.0, "Agricole territorial": 0.8, "Industriel territorial": 0.5}.get(t, 0)


def priorite(score, base):
    p = PRIO[base]
    return "A" if score >= p["A"] else "B" if score >= p["B"] else "C" if score >= p["C"] else "D"


# ---------------------------------------------------------- règles partagées
def cpb_coef(annee_mes):
    c = PARAMS["cpb"]
    if not annee_mes:
        return None, None, None
    butoir = int(c["date_butoir_injection"][:4])
    conv = c["annee_conversion_defaut"]
    age = conv - annee_mes
    coef = 1 if age < c["age_min_ans"] else (c["coefficient_majore"] if age <= c["age_max_ans"] and conv <= butoir else c["coefficient_base"])
    first, last = max(ANNEE, annee_mes + c["age_min_ans"]), min(butoir, annee_mes + c["age_max_ans"])
    return coef, first, last


def score_elec(d):
    w = W["elec"]
    kw = d.get("puissance_kw")
    mes = d.get("annee_mes")
    ech = mes + PARAMS["cogen"]["duree_contrat_ans"] if mes else None
    coef, first, last = cpb_coef(mes)
    fc = (d["energie_gwh_an"] * 1000 / (kw * 8.76)) if d.get("energie_gwh_an") and kw else None
    detail = {
        "puissance": round(w["puissance"] * frac_puissance(kw), 1),
        "fenetre_echeance": round(w["fenetre_echeance"] * frac_echeance(ech), 1),
        "coefficient_cpb": round(w["coefficient_cpb"] * frac_coef(coef), 1),
        "reseau": round(w["reseau"] * frac_reseau(d.get("dist_grdf_km")), 1),
        "facteur_charge": round(w["facteur_charge"] * frac_charge(fc), 1),
    }
    s = round(sum(detail.values()), 1)
    return s, detail, {"echeance": ech, "coef": coef, "fenetre": (first, last) if first is not None and first <= last else None, "fc": round(fc, 2) if fc is not None else None}


def score_inj(d):
    w = W["injection"]
    cap = d.get("capacite_de_production_gwh_an")
    mes = d.get("annee_mes")
    rest = (mes + PARAMS["injection"]["duree_tarif_ans"] - ANNEE) if mes else None
    detail = {
        "capacite": round(w["capacite"] * frac_capacite(cap), 1),
        "tarif_residuel": round(w["tarif_residuel"] * frac_tarif(rest), 1),
        "type": round(w["type"] * frac_type(d.get("site")), 1),
        "augmentation_prevue": round(w["augmentation_prevue"] * (1.0 if (d.get("augmentation_prevue") or "").startswith("Augmentation") else 0.4), 1),
    }
    return round(sum(detail.values()), 1), detail, {"tarif_restant": rest}


# ------------------------------------------------------------------- Excel
NAVY, TEAL, PALE = "1E4260", "22788C", "F3F6F9"


def sheet_table(wb, title, headers, rows, widths=None, freeze="C2"):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", name="Arial", size=9)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=9)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 42
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(38, len(h) + 2)))
    return ws


ELEC_HEADERS = ["Rang", "Priorité", "Score /100", "Pts puissance", "Pts échéance", "Pts coef. CPB", "Pts réseau", "Pts facteur charge",
                "Nom", "Commune", "Département", "Région", "Zone test", "Puissance kWé", "Énergie injectée GWh/an", "Facteur de charge",
                "Technologie", "Mise en service", "Échéance estimée", "Coef. CPB 2028", "Fenêtre 0,95", "Distance GRDF km",
                "Injection la plus proche km", "Injection la plus proche", "Zonage", "Zonage maturité", "Zonage capacité max Nm3/h",
                "Gestionnaire", "Code EIC", "Latitude", "Longitude", "Google Maps",
                "Contact", "Statut prospection", "Commentaires AdlF"]
INJ_HEADERS = ["Rang", "Priorité", "Score /100", "Pts capacité", "Pts tarif résiduel", "Pts type", "Pts augmentation",
               "Nom", "Commune", "Département", "Région", "Zone test", "Type de site", "Capacité GWh/an", "Mise en service",
               "Tarif restant (ans)", "Échéance tarif", "Augmentation prévue", "Réseau", "Gestionnaire", "PITD/PITP",
               "Zonage", "Zonage maturité", "Latitude", "Longitude", "Google Maps",
               "Contact", "Statut prospection", "Commentaires AdlF"]


def elec_row(rank, d):
    z = d.get("zonage") or {}
    f = d["score_meta"]["fenetre"]
    lat, lon = d.get("lat"), d.get("lon")
    return [rank, d["priorite_v2"], d["score_v2"], d["score_detail"]["puissance"], d["score_detail"]["fenetre_echeance"],
            d["score_detail"]["coefficient_cpb"], d["score_detail"]["reseau"], d["score_detail"]["facteur_charge"],
            d["nom"], d["commune"], d["departement"], d["region"], "oui" if d["region"] in PARAMS["geographie"]["zone_test"] else "",
            d.get("puissance_kw"), d.get("energie_gwh_an"), d["score_meta"]["fc"], d.get("technologie"), d.get("date_mes"),
            d["score_meta"]["echeance"], d["score_meta"]["coef"], f"{f[0]}-{f[1]}" if f else "hors d'atteinte",
            d.get("dist_grdf_km"), d.get("dist_injection_km"), d.get("injection_proche"),
            z.get("libelle"), z.get("maturite"), z.get("capamax"), d.get("gestionnaire"), d.get("code_eic"), lat, lon,
            f"https://www.google.com/maps?q={lat},{lon}" if lat is not None else "", "", "", ""]


def inj_row(rank, d):
    z = d.get("zonage") or {}
    c = d.get("coordonnees") or {}
    return [rank, d["priorite_v2"], d["score_v2"], d["score_detail"]["capacite"], d["score_detail"]["tarif_residuel"],
            d["score_detail"]["type"], d["score_detail"]["augmentation_prevue"],
            d.get("nom_du_projet"), d.get("commune"), d.get("departement"), d.get("region"),
            "oui" if d.get("region") in PARAMS["geographie"]["zone_test"] else "", d.get("site"),
            d.get("capacite_de_production_gwh_an"), d.get("date_de_mes"), d["score_meta"]["tarif_restant"],
            (d.get("annee_mes") or 0) + PARAMS["injection"]["duree_tarif_ans"] if d.get("annee_mes") else None,
            d.get("augmentation_prevue"), d.get("type_de_reseau"), d.get("grx_demandeur"), d.get("ndeg_de_pitd_pitp"),
            z.get("libelle"), z.get("maturite"), c.get("lat"), c.get("lon"),
            f"https://www.google.com/maps?q={c.get('lat')},{c.get('lon')}" if c else "", "", "", ""]


def methodo_sheet(wb, n_elec, n_inj, meta):
    ws = wb.active
    ws.title = "Méthodologie"
    lines = [
        ("Nautilus — Qualification des sites biométhane France, score v2", True),
        (f"Généré le {date.today().isoformat()} par tools/qualify_v2.py (dépôt biomethane-france). Statut : proposition de pondération, à valider avec Antoine de la Faire.", False),
        ("", False),
        ("Sources", True),
        (f"Points d'injection : ODRÉ, {meta['injection']['source']} ({meta['injection']['extraction']}).", False),
        (f"Électricité biogaz : ODRÉ, {meta['cogen']['source']} (traitement {meta['cogen'].get('data_processed')}).", False),
        ("Réseau : cartographie du réseau GRDF en service (open data GRDF), zonages de raccordement (ODRÉ, 2020), distance au point d'injection le plus proche.", False),
        ("Règles : arrêté du 10/08/2026 (guichet), arrêté du 26/12/2025 (coefficient CPB), enseignements BC 14/09/2026, reprise AdlF 18/09/2026, annexe réglementaire 24/09/2026.", False),
        ("", False),
        (f"Périmètre scoré — électricité biogaz ({n_elec} sites) : code combustible B.MET (méthanisation, toutes technologies), en service, puissance ≥ {PARAMS['cogen']['puissance_kw']['plancher']} kWé.", False),
        (f"Périmètre scoré — injection ({n_inj} sites) : {', '.join(PARAMS['injection']['types'])}, site ouvert, ≥ {PARAMS['injection']['capacite_gwh_an']['min']} GWh/an.", False),
        ("", False),
        ("Pondération électricité biogaz (total 100)", True),
    ]
    for k, v in W["elec"].items():
        lines.append((f"{k} : {v} points — {GRILLES['elec'][k]}", False))
    lines += [("", False), ("Pondération injection (total 100)", True)]
    for k, v in W["injection"].items():
        lines.append((f"{k} : {v} points — {GRILLES['injection'][k]}", False))
    lines += [("", False),
              (f"Priorités électricité biogaz : A ≥ {PRIO['elec']['A']} · B ≥ {PRIO['elec']['B']} · C ≥ {PRIO['elec']['C']} · D en dessous. "
               f"Injection : A ≥ {PRIO['injection']['A']} · B ≥ {PRIO['injection']['B']} · C ≥ {PRIO['injection']['C']} (seuils calibrés pour A ≈ 13 % de chaque base).", True),
              ("", False),
              ("Limites", True),
              ("Positions des installations électriques au centroïde de commune (± quelques km) ; distance GRDF à vol d'oiseau, réseau GRDF seulement (zones ELD = inconnue).", False),
              ("Échéances estimées (MES + 20 ans BG, MES + 15 ans injection), avenants non captés. Coefficient CPB estimé à l'année de conversion par défaut (2028).", False),
              ("Pas de donnée d'intrants, de capital ni de gouvernance : le facteur de charge tient lieu de proxy d'exploitation en attendant le registre équipe (étape 8) et Pappers (étape 9).", False),
              ("Ce score trie ; il ne remplace ni l'étude détaillée GRDF ni l'entretien avec l'exploitant.", False),
              ("", False),
              ("Colonnes « Contact », « Statut prospection », « Commentaires AdlF » laissées vides pour la qualification terrain.", False)]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=2, value=t)
        c.font = Font(name="Arial", size=10 if not bold else 11, bold=bold, color=NAVY if bold else "1A1A1A")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 140


def main():
    xlsx = DEFAULT_XLSX
    if "--xlsx" in sys.argv:
        xlsx = sys.argv[sys.argv.index("--xlsx") + 1]
    inj = json.load(open(DATA / "points-injection.json", encoding="utf-8"))
    elec = json.load(open(DATA / "cogenerations.json", encoding="utf-8"))
    meta = json.load(open(DATA / "meta.json", encoding="utf-8"))
    pl = PARAMS["cogen"]["puissance_kw"]["plancher"]

    for d in elec:
        ok = d.get("code_combustible") == "B.MET" and d.get("statut") == "En service" and (d.get("puissance_kw") or 0) >= pl
        if ok:
            s, det, m = score_elec(d)
            d["score_v2"], d["score_detail"], d["priorite_v2"], d["score_meta"] = s, det, priorite(s, "elec"), m
        else:
            d["score_v2"] = d["priorite_v2"] = d["score_detail"] = None
            d.pop("score_meta", None)
    for d in inj:
        cmin = PARAMS["injection"]["capacite_gwh_an"]["min"] or 0
        ok = d.get("site") in PARAMS["injection"]["types"] and d.get("site_ouvert") == "True" and (d.get("capacite_de_production_gwh_an") or 0) >= cmin
        if ok:
            s, det, m = score_inj(d)
            d["score_v2"], d["score_detail"], d["priorite_v2"], d["score_meta"] = s, det, priorite(s, "injection"), m
        else:
            d["score_v2"] = d["priorite_v2"] = d["score_detail"] = None
            d.pop("score_meta", None)

    se = sorted([d for d in elec if d["score_v2"] is not None], key=lambda d: (-d["score_v2"], -(d.get("puissance_kw") or 0)))
    si = sorted([d for d in inj if d["score_v2"] is not None], key=lambda d: (-d["score_v2"], -(d.get("capacite_de_production_gwh_an") or 0)))
    zt = [d for d in se if d["region"] in PARAMS["geographie"]["zone_test"]]

    wb = Workbook()
    methodo_sheet(wb, len(se), len(si), meta)
    sheet_table(wb, "Élec biogaz — Top 50", ELEC_HEADERS, [elec_row(i + 1, d) for i, d in enumerate(se[:50])])
    sheet_table(wb, "Injection — Top 50", INJ_HEADERS, [inj_row(i + 1, d) for i, d in enumerate(si[:50])])
    sheet_table(wb, "Zone test — Top 30 élec", ELEC_HEADERS, [elec_row(i + 1, d) for i, d in enumerate(zt[:30])])
    sheet_table(wb, f"Élec biogaz — {len(se)} scorés", ELEC_HEADERS, [elec_row(i + 1, d) for i, d in enumerate(se)])
    sheet_table(wb, f"Injection — {len(si)} scorés", INJ_HEADERS, [inj_row(i + 1, d) for i, d in enumerate(si)])
    Path(xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx)

    # données de l'app : score, priorité, détail (score_meta reste interne au classeur)
    for d in elec + inj:
        d.pop("score_meta", None)
    json.dump(elec, open(DATA / "cogenerations.json", "w", encoding="utf-8"), ensure_ascii=False)
    json.dump(inj, open(DATA / "points-injection.json", "w", encoding="utf-8"), ensure_ascii=False)
    meta["score_v2"] = {"date": date.today().isoformat(), "ponderation": W, "xlsx": str(xlsx)}
    json.dump(meta, open(DATA / "meta.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)

    from collections import Counter
    print(f"Élec biogaz scorés : {len(se)} — priorités {dict(Counter(d['priorite_v2'] for d in se))} — "
          f"zone test {len(zt)} ({dict(Counter(d['priorite_v2'] for d in zt))})")
    print(f"Injection scorés : {len(si)} — priorités {dict(Counter(d['priorite_v2'] for d in si))}")
    print("Top 5 élec :", [(d['nom'], d['commune'], d['score_v2']) for d in se[:5]])
    print("Top 5 injection :", [(d['nom_du_projet'], d['commune'], d['score_v2']) for d in si[:5]])
    print("Classeur :", xlsx)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
